"""Turn uploaded files into a list of pages of plain text."""

from __future__ import annotations

import io
import logging
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from .text_utils import normalize_text

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFileType(ValueError):
    pass


@dataclass
class Page:
    number: int  # 1-based
    text: str


@dataclass
class LoadedDocument:
    pages: list[Page]
    warnings: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)


def load_document(filename: str, data: bytes, enable_ocr: bool = False, tesseract_cmd: str = "") -> LoadedDocument:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        doc = _load_pdf(data, enable_ocr, tesseract_cmd)
    elif ext == ".docx":
        doc = _load_docx(data)
    elif ext in {".txt", ".md"}:
        doc = LoadedDocument([Page(1, _decode_text(data))])
    else:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    doc.pages = [Page(p.number, normalize_text(p.text)) for p in doc.pages]
    doc.pages = [p for p in doc.pages if p.text]
    if not doc.pages:
        doc.warnings.append(
            "No text could be extracted. If this is a scanned PDF, set ENABLE_OCR=true "
            "and install Tesseract with Hindi (hin) language data."
        )
    return doc


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


# Devanagari dependent signs: candrabindu/anusvara/visarga, nukta, vowel signs (matras), virama.
_DEVA_SIGNS = "\u0900-\u0903\u093a-\u094f\u0955-\u0957\u0962\u0963"
# The same vowel sign twice within one syllable, possibly with a nasal/nukta/virama between
# ('ाा', 'ांा'): never valid Hindi, but common in PDF text layers.
_DUPLICATE_SIGN = re.compile("([\u093e-\u094c\u0962\u0963])([\u0900-\u0903\u093c\u094d]*)\\1+")
_DUPLICATE_MARK = re.compile("([ऀ-ः़])\\1+")  # doubled anusvara etc.: 'ंं' -> 'ं'
_ORPHAN_SIGN = re.compile(f"(?:^|\\s)[{_DEVA_SIGNS}]", re.MULTILINE)
_DEVA_WORD = re.compile("\\S*[\u0900-\u097f]\\S*")


def clean_devanagari_extraction(text: str) -> str:
    """PDF text layers often repeat a matra (e.g. 'महाविद्याालय'); valid Hindi never has
    the same vowel sign twice in one syllable, so collapsing repeats is safe."""
    return _DUPLICATE_MARK.sub(r"\1", _DUPLICATE_SIGN.sub(r"\1\2", text))


def looks_garbled_hindi(text: str) -> bool:
    """True if a PDF's Hindi text layer is broken.

    Many Hindi PDFs store glyphs in *visual* order and drop conjuncts, so extraction gives
    'िनयम' instead of 'नियम': a vowel sign at the start of a word, which is impossible in
    valid Unicode Hindi. When that is common, the page should be OCR'd instead.
    """
    deva_words = len(_DEVA_WORD.findall(text))
    if deva_words < 10:
        return False
    return len(_ORPHAN_SIGN.findall(text)) / deva_words > 0.05


def _extract_pdf_pages(data: bytes) -> list[str]:
    """Per-page text. PyMuPDF (MuPDF) handles Devanagari far better than pypdf, which drops
    conjuncts and reorders matras, so it is preferred when installed."""
    try:
        import pymupdf
    except ImportError:
        pymupdf = None
    if pymupdf is not None:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return [_mupdf_page_text(page) for page in doc]

    from pypdf import PdfReader

    texts = []
    for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
        try:
            texts.append(page.extract_text() or "")
        except Exception as exc:  # malformed page objects are common in the wild
            log.warning("Failed to extract page %d: %s", i, exc)
            texts.append("")
    return texts


def _mupdf_page_text(page) -> str:
    """Rebuild lines from MuPDF's layout.

    MuPDF sometimes splits one visual line in the middle of a Devanagari word (around the
    'ि' matra, which is drawn to the left of its consonant), giving 'उपस्थि' + 'ति'.
    Fragments on the same baseline are re-joined: directly when both sides of the join
    are Devanagari with no space character between them, otherwise with a space when
    there is a visible gap (e.g. between table cells). Consecutive blocks are separated
    by a blank line only when there is a real vertical gap (a new paragraph).
    """
    rows: list[list] = []  # [baseline_y, x_end, font_size, text, block_top]
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:  # skip images
            continue
        for line in block["lines"]:
            text = "".join(s["text"] for s in line["spans"])
            if not text.strip():
                continue
            x0, y0, x1, y1 = line["bbox"]
            size = max((s["size"] for s in line["spans"]), default=10)
            prev = rows[-1] if rows else None
            if prev and abs(prev[0] - y1) < 0.3 * size and x0 >= prev[1] - 2 * size:
                left, right = prev[3], text
                gap = x0 - prev[1]
                if left[-1:].isspace() or right[:1].isspace():
                    joiner = ""
                elif _is_devanagari(left[-1]) and _is_devanagari(right[0]) and gap < 1.5 * size:
                    joiner = ""  # split inside a word; the reordered matra inflates the gap
                else:
                    joiner = " " if gap > 0.25 * size else ""
                prev[3] = left + joiner + right
                prev[1] = max(prev[1], x1)
            else:
                rows.append([y1, x1, size, text, y0])

    out: list[str] = []
    for i, row in enumerate(rows):
        if i:
            prev = rows[i - 1]
            paragraph_gap = row[4] - prev[0] > 0.9 * prev[2]
            out.append("\n\n" if paragraph_gap else "\n")
        out.append(row[3].strip())
    return "".join(out)


def _is_devanagari(ch: str) -> bool:
    return "ऀ" <= ch <= "ॿ"


def _load_pdf(data: bytes, enable_ocr: bool, tesseract_cmd: str = "") -> LoadedDocument:
    pages = [Page(i, clean_devanagari_extraction(t)) for i, t in enumerate(_extract_pdf_pages(data), start=1)]
    empty = [p.number for p in pages if len(p.text.strip()) < 20]
    broken = [p.number for p in pages if p.number not in empty
              and (looks_garbled_hindi(p.text) or _looks_like_legacy_hindi_font(p.text))]

    warnings: list[str] = []
    needs_ocr = empty + broken
    if needs_ocr and enable_ocr:
        ocr_text, error = _ocr_pdf_pages(data, needs_ocr, tesseract_cmd)
        for page in pages:
            if page.number in ocr_text:
                page.text = ocr_text[page.number]
        if ocr_text:
            reason = " (scanned or garbled Hindi text layer)" if broken else ""
            warnings.append(f"OCR (Tesseract hin+eng) was used for {len(ocr_text)} page(s){reason}.")
        if error:
            warnings.append(error)
    else:
        if empty:
            warnings.append(
                f"{len(empty)} page(s) had no extractable text (scanned images?). Enable OCR to index them."
            )
        if broken:
            warnings.append(
                f"{len(broken)} page(s) have a garbled Hindi text layer (legacy font or broken PDF "
                "encoding). Enable OCR for accurate text."
            )
    return LoadedDocument(pages, warnings)


# The most frequent Hindi words as they appear when typed in Kruti Dev (a pre-Unicode font
# that draws Devanagari glyphs on Latin code points): का के की में है और कि को से पर हेतु
# द्वारा लिए तथा एवं नाम. None of these are English words.
_KRUTI_DEV_WORDS = {
    "dk", "ds", "dh", "esa", "gS", "gSa", "vkSj", "fd", "dks", "ls", "ij", "gsrq",
    "}kjk", "fy,", "rFkk", ",oa", "uke", "fd;k", "tk,xkA", "gksxkA", "dj", "Hkh",
}
# Symbols that Kruti Dev uses *inside* words for glyphs: 'fo|ky;' (विद्यालय), 'uke%&' (नाम:-).
_KRUTI_DEV_INNER = re.compile(r"[A-Za-z][;|}{\]\[~`][A-Za-z]|%&|[¼½]")


def _looks_like_legacy_hindi_font(text: str) -> bool:
    """Detect Kruti Dev / DevLys-style text: Hindi typed with a legacy font, which extracts
    as Latin gibberish such as 'fo|ky; dk uke%&' (विद्यालय का नाम:-)."""
    words = re.findall(r"\S+", text[:8000])
    if len(words) < 20:
        return False
    common = sum(w.strip(".,:()\"'") in _KRUTI_DEV_WORDS for w in words)
    inner = len(_KRUTI_DEV_INNER.findall(text[:8000]))
    return (common + inner) / len(words) > 0.06


_DEFAULT_WINDOWS_TESSERACT = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")


def find_tesseract(configured: str = "") -> str | None:
    """TESSERACT_CMD if set, else `tesseract` on PATH, else the default Windows install."""
    if configured:
        return configured if Path(configured).exists() else None
    on_path = shutil.which("tesseract")
    if on_path:
        return on_path
    return str(_DEFAULT_WINDOWS_TESSERACT) if _DEFAULT_WINDOWS_TESSERACT.exists() else None


def _ocr_pdf_pages(data: bytes, page_numbers: list[int], tesseract_cmd: str = "") -> tuple[dict[int, str], str | None]:
    """OCR the given pages. Returns ({page_number: text}, error message or None)."""
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError:
        return {}, "OCR is enabled but pytesseract/pypdfium2 are not installed (pip install pytesseract pypdfium2)."
    cmd = find_tesseract(tesseract_cmd)
    if cmd is None:
        return {}, "OCR is enabled but tesseract.exe was not found. Install Tesseract or set TESSERACT_CMD."
    pytesseract.pytesseract.tesseract_cmd = cmd

    results: dict[int, str] = {}
    failed = 0
    pdf = pdfium.PdfDocument(data)
    for number in page_numbers:
        try:
            # Render at ~180 DPI (2.5 x 72); Tesseract accuracy drops sharply below ~150 DPI.
            image = pdf[number - 1].render(scale=2.5).to_pil()
            results[number] = pytesseract.image_to_string(image, lang="hin+eng")
        except Exception as exc:
            failed += 1
            log.warning("OCR failed on page %d: %s", number, exc)
    return results, (f"OCR failed on {failed} page(s); see server log." if failed else None)


def _load_docx(data: bytes) -> LoadedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    # DOCX has no fixed pages; treat the whole document as page 1.
    return LoadedDocument([Page(1, "\n".join(parts))])
